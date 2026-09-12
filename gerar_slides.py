import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import comtypes.client

def main():
    prs = Presentation()
    
    # Remove default slides if we want to add from blank, or just use blank
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    BG_COLOR = RGBColor(22, 30, 45)
    TITLE_COLOR = RGBColor(255, 192, 0)
    TEXT_COLOR = RGBColor(255, 255, 255)
    
    def set_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR
        
    def add_title(slide, text):
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.bold = True
        p.font.size = Pt(44)
        p.font.color.rgb = TITLE_COLOR
        return tb

    # 1. TITLE
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s1)
    tb = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "FRIEDRICH NIETZSCHE"
    p.font.bold = True
    p.font.size = Pt(60)
    p.font.color.rgb = TEXT_COLOR
    p.alignment = PP_ALIGN.CENTER
    p2 = tb.text_frame.add_paragraph()
    p2.text = "O Filósofo do Martelo (1844 – 1900)"
    p2.font.size = Pt(28)
    p2.font.color.rgb = RGBColor(200, 200, 200)
    p2.alignment = PP_ALIGN.CENTER
    
    sh = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(4.5), Inches(9.33), Inches(1.5))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(15, 20, 30)
    sh.line.color.rgb = RGBColor(50, 100, 150)
    p3 = sh.text_frame.paragraphs[0]
    p3.text = '"E se a vida que você vive hoje se repetisse idêntica pela eternidade?\nVocê a viveria com alegria ou desespero?"'
    p3.font.size = Pt(24)
    p3.font.italic = True
    p3.font.color.rgb = RGBColor(100, 180, 255)
    p3.alignment = PP_ALIGN.CENTER

    # 2. HOMEM POR TRÁS DO MARTELO
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s2)
    add_title(s2, "O Homem por Trás do Martelo")
    tb2 = s2.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(12.33), Inches(5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    bullets = [
        "Contexto Histórico: Século XIX. O apogeu da ciência e da razão na Europa e o declínio vertiginoso da religião tradicional.",
        "Prodígio Intelectual: Nascido na Prússia, foi um gênio precoce. Tornou-se professor universitário de filologia (estudo de textos antigos) aos 24 anos.",
        "A Vida Errante: Devido à saúde extremamente frágil (fortes enxaquecas e problemas de visão), abandonou a academia e viveu viajando pelos Alpes, escrevendo de forma intensa.",
        "O Fim Trágico: Sofreu um colapso mental severo em 1889 nas ruas de Turim, Itália. Passou seus últimos 11 anos alienado, cuidado por sua irmã (que mais tarde distorceu seus textos)."
    ]
    for i, b in enumerate(bullets):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(22)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(20)

    # 3. O ABISMO DO NIILISMO
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s3)
    add_title(s3, "O Abismo do Niilismo")
    tb3 = s3.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.33), Inches(1))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = 'O que acontece com a humanidade quando as velhas "Verdades Absolutas" desaparecem?'
    p3.font.size = Pt(28)
    p3.font.color.rgb = TEXT_COLOR
    p3.alignment = PP_ALIGN.CENTER
    
    sh3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.8), Inches(10.33), Inches(3.2))
    sh3.fill.solid()
    sh3.fill.fore_color.rgb = RGBColor(30, 40, 60)
    sh3.line.color.rgb = RGBColor(255, 50, 50)
    tf3 = sh3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "O Diagnóstico de uma Era"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = TITLE_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    p = tf3.add_paragraph()
    p.text = "Com o avanço da ciência e do iluminismo, a base que sustentava a moralidade humana (Deus, as religiões, a tradição) ruiu. Nietzsche percebeu que isso geraria uma crise sem precedentes."
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR
    p.space_before = Pt(14)
    
    p = tf3.add_paragraph()
    p.text = 'O perigo iminente era o NIILISMO (do latim nihil, "nada"): a crença assustadora de que a vida não tem nenhum sentido, nenhum valor intrínseco e nenhum propósito superior.'
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR
    p.space_before = Pt(14)
    
    tb3f = s3.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11.33), Inches(0.8))
    p = tb3f.text_frame.paragraphs[0]
    p.text = "Como o ser humano pode evitar o desespero e encontrar força para viver em um universo vazio de sentidos prontos?"
    p.font.size = Pt(22)
    p.font.italic = True
    p.font.color.rgb = RGBColor(100, 180, 255)
    p.alignment = PP_ALIGN.CENTER

    # 4. TRANSVALORAÇÃO
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s4)
    add_title(s4, "A Transvaloração dos Valores")
    tb4 = s4.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.33), Inches(1))
    p = tb4.text_frame.paragraphs[0]
    p.text = "A resposta de Nietzsche para o abismo do Niilismo é radical. Se não existem valores absolutos flutuando no universo esperando para serem descobertos, nós devemos criá-los."
    p.font.size = Pt(24)
    p.font.color.rgb = TEXT_COLOR
    
    sh4_1 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2.8), Inches(5.8), Inches(4))
    sh4_1.fill.solid()
    sh4_1.fill.fore_color.rgb = RGBColor(22, 30, 45)
    sh4_1.line.fill.background()
    tf = sh4_1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🚫 O Problema Atual"
    p.font.bold = True
    p.font.size = Pt(26)
    p.font.color.rgb = TITLE_COLOR
    p = tf.add_paragraph()
    p.text = 'Vivemos sob uma "Moral de Rebanho" (herança cultural europeia) que reprime os instintos, valoriza a fraqueza, a obediência e promete recompensas apenas em uma "vida após a morte", negando a vida real.'
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR
    
    sh4_2 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(2.8), Inches(5.8), Inches(4))
    sh4_2.fill.solid()
    sh4_2.fill.fore_color.rgb = RGBColor(22, 30, 45)
    sh4_2.line.fill.background()
    tf = sh4_2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔨 A Solução Proposta"
    p.font.bold = True
    p.font.size = Pt(26)
    p.font.color.rgb = TITLE_COLOR
    p = tf.add_paragraph()
    p.text = 'É preciso filosofar "com o martelo": destruir as velhas ilusões morais e realizar a Transvaloração. O indivíduo não deve buscar ser "bom" segundo o rebanho, mas deve afirmar a vida em sua totalidade, aceitando a dor e o prazer, tornando-se o autor do seu próprio destino.'
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR

    # 5. VOCABULÁRIO NIETZSCHEANO
    s5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s5)
    add_title(s5, "Vocabulário Nietzscheano")
    
    vocabs = [
        ("1. A Morte de Deus", 'Não é uma comemoração de ateus, mas a constatação de um evento cultural. A "Morte de Deus" significa o fim das garantias absolutas. Não há mais um fundamento inquestionável para a moralidade e para a verdade.', Inches(0.5), Inches(1.5)),
        ("2. Vontade de Potência", "A força motriz essencial de tudo que vive. Não é o desejo de poder político para dominar os outros, mas o instinto vital de crescer, expandir-se, criar, superar obstáculos e se tornar mais forte.", Inches(6.8), Inches(1.5)),
        ("3. Eterno Retorno & Amor Fati", "O teste supremo: viver de tal modo que você desejaria que sua vida se repetisse eternamente. O Amor Fati (Amor ao Destino) é abraçar alegremente a vida exatamente como ela é, com suas dores e glórias.", Inches(0.5), Inches(4.5)),
        ("4. O Übermensch", 'Traduzido como "Além-do-Homem" ou "Super-Homem". É o indivíduo livre que superou a moral do rebanho, não sente ressentimento pela vida, cria seus próprios valores e vive guiado por sua vontade de potência.', Inches(6.8), Inches(4.5))
    ]
    
    for t, d, l, t_pos in vocabs:
        sh = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t_pos, Inches(6), Inches(2.6))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(30, 40, 60)
        sh.line.color.rgb = RGBColor(50, 100, 150)
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_COLOR
        p = tf.add_paragraph()
        p.text = d
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(220, 220, 220)

    # 6. O PROCESSO DE SUPERAÇÃO
    s6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s6)
    tb = s6.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.33), Inches(1))
    p = tb.text_frame.paragraphs[0]
    p.text = "O Processo de Superação do Espírito"
    p.font.bold = True
    p.font.size = Pt(40)
    p.font.color.rgb = TITLE_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    tb_sub = s6.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.33), Inches(1))
    p = tb_sub.text_frame.paragraphs[0]
    p.text = 'Em sua obra Assim Falava Zaratustra, Nietzsche utiliza três metáforas brilhantes para explicar como o ser humano abandona o "rebanho" para se tornar criador de si mesmo.'
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    steps = [
        ("1. O Camelo", '"Tu deves!"\nO espírito de carga. É forte, mas submisso. Aceita carregar as tradições, a moral da sociedade, os medos e as obrigações impostas pelos outros. Suporta o peso do passado no deserto do niilismo.', Inches(0.5), RGBColor(255, 192, 0)),
        ("2. O Leão", '"Eu quero!"\nO espírito da revolta. O camelo se transforma em leão no deserto para enfrentar o dragão dos velhos valores. Ele destrói a submissão, conquista a liberdade e o espaço para a novidade, mas ainda não sabe criar.', Inches(4.7), RGBColor(255, 50, 50)),
        ("3. A Criança", '"Eu sou!"\nO espírito criador (O Übermensch). A criança é a inocência, o esquecimento das velhas regras, um recomeço sagrado. É o jogo criativo de inventar a própria vida e seus novos valores com alegria.', Inches(8.9), RGBColor(50, 150, 255))
    ]
    for title, desc, left, col in steps:
        sh = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.8), Inches(4), Inches(4.2))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(30, 40, 60)
        sh.line.color.rgb = col
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_COLOR
        p.alignment = PP_ALIGN.CENTER
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(220, 220, 220)
        p.alignment = PP_ALIGN.CENTER

    # 7. NA PRÁTICA
    s7 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s7)
    add_title(s7, "Na Prática: A Escolha do Futuro")
    tb7 = s7.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.33), Inches(1))
    p = tb7.text_frame.paragraphs[0]
    p.text = "A Situação: Rafael, 17 anos, gosta de artes e tecnologia, mas sofre forte pressão para cursar Direito, pois é o que sua família considera \"seguro e respeitável\"."
    p.font.size = Pt(22)
    p.font.color.rgb = TEXT_COLOR
    
    sh7_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.5), Inches(12.33), Inches(1.5))
    sh7_1.fill.solid()
    sh7_1.fill.fore_color.rgb = RGBColor(30, 40, 60)
    tf = sh7_1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👥 A Ótica do Rebanho\nRafael cede à pressão por medo do julgamento. Escolhe o caminho seguro para não causar conflitos. Foge do sofrimento da desaprovação, mas anula seus próprios instintos e vive uma vida burocrática e ressentida."
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT_COLOR
    
    sh7_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(4.2), Inches(12.33), Inches(1.5))
    sh7_2.fill.solid()
    sh7_2.fill.fore_color.rgb = RGBColor(30, 40, 60)
    tf = sh7_2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ A Ótica Nietzscheana (Vontade de Potência)\nRafael se torna o \"Leão\", rompe com a imposição externa (\"Tu Deves\") e assume o risco de sua escolha autêntica. Ele entende que a verdadeira grandeza exige suportar a dor do conflito para esculpir sua própria obra (sua vida)."
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT_COLOR
    
    tb7f = s7.shapes.add_textbox(Inches(0.5), Inches(6), Inches(12.33), Inches(1))
    p = tb7f.text_frame.paragraphs[0]
    p.text = "Conclusão: A vida autêntica não é a mais pacífica. Ela exige coragem para decepcionar os outros em nome de se tornar quem você realmente é."
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(100, 180, 255)

    # 8. DIÁLOGO E CONTROVÉRSIAS
    s8 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s8)
    add_title(s8, "Diálogo e Controvérsias")
    
    sh8_1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(6), Inches(5.5))
    sh8_1.fill.solid()
    sh8_1.fill.fore_color.rgb = RGBColor(30, 40, 60)
    tf = sh8_1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "O Choque: Kant vs. Nietzsche"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = TEXT_COLOR
    p = tf.add_paragraph()
    p.text = "Dois gigantes da filosofia alemã com visões opostas sobre como devemos agir.\n\nPonto em Comum: Ambos defendem a autonomia humana (guiar a si mesmo, sair da menoridade).\n\nA Grande Divergência:\n• Immanuel Kant: Defende uma regra universal e racional para todos (O Imperativo Categórico). A moral é dever e obediência à razão.\n• Friedrich Nietzsche: Rejeita qualquer regra universal. Uma moral que serve para todos é moral de rebanho. Cada indivíduo forte deve criar sua própria lei."
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT_COLOR

    sh8_2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(6), Inches(5.5))
    sh8_2.fill.solid()
    sh8_2.fill.fore_color.rgb = RGBColor(30, 40, 60)
    tf = sh8_2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Principais Críticas e Mitos"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = TEXT_COLOR
    p = tf.add_paragraph()
    p.text = "⚠️ O Mito do Nazismo\nNietzsche não era nazista ou antissemita (ele repudiava fortemente o nacionalismo). Após sua morte, sua irmã, Elisabeth, filiada ao nazismo, falsificou e publicou trechos de seus diários para agradar ao Terceiro Reich.\n\n🔨 Críticas Acadêmicas\n• Elitismo extremo: Seu profundo desdém pelas massas (o \"rebanho\") torna difícil aplicar suas ideias para construir uma sociedade democrática.\n• Falta de projeto coletivo: É um pensador focado na salvação individual. Não oferece soluções para a injustiça social sistêmica."
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT_COLOR

    # 9. POR QUE LER NIETZSCHE HOJE
    s9 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s9)
    add_title(s9, "Por que ler Nietzsche Hoje?")
    
    sh9_1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(6), Inches(5.5))
    sh9_1.fill.solid()
    sh9_1.fill.fore_color.rgb = RGBColor(30, 40, 60)
    tf = sh9_1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Impacto Histórico"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = TITLE_COLOR
    p = tf.add_paragraph()
    p.text = "• Psicanálise: Antecipou as descobertas de Freud e Jung sobre a força do inconsciente e dos instintos reprimidos na mente humana.\n\n• Existencialismo: Abriu as portas para Sartre e Camus, ao declarar que a vida não tem um sentido pré-determinado, obrigando o homem a ser livre e responsável.\n\n• Filosofia Contemporânea: Inspirou pensadores como Foucault a investigar como a \"Verdade\" é muitas vezes apenas uma máscara para o poder."
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR

    sh9_2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(6), Inches(5.5))
    sh9_2.fill.solid()
    sh9_2.fill.fore_color.rgb = RGBColor(22, 30, 45) # Different bg
    sh9_2.line.color.rgb = TITLE_COLOR
    tf = sh9_2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "O Antídoto Moderno"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = TITLE_COLOR
    p = tf.add_paragraph()
    p.text = "Nietzsche é incrivelmente atual para o jovem do século XXI porque atua como um choque contra os males modernos:\n\n✓ Contra a positividade tóxica: Ensina que o sofrimento não deve ser escondido, mas transformado em ferramenta de grandeza (Amor Fati).\n\n✓ Contra as redes sociais: Denuncia a mentalidade de \"rebanho\" dos algoritmos, os \"cancelamentos\" por medo da maioria e a busca desesperada por aceitação.\n\n✓ Contra o vazio: Exige que abandonemos a apatia e assumamos a postura de artistas criando a grande obra que é a nossa própria vida."
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR

    # 10. TESTE SUA COMPREENSÃO
    s10 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s10)
    
    # Left column concepts
    sh10_1 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.5), Inches(4), Inches(1))
    sh10_1.fill.solid(); sh10_1.fill.fore_color.rgb = RGBColor(30, 40, 60)
    p = sh10_1.text_frame.paragraphs[0]
    p.text = "1 FRASE DE OURO\n\"Torna-te quem tu és.\""
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_COLOR
    
    sh10_2 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.7), Inches(4), Inches(2))
    sh10_2.fill.solid(); sh10_2.fill.fore_color.rgb = RGBColor(30, 40, 60)
    p = sh10_2.text_frame.paragraphs[0]
    p.text = "3 IDEIAS FUNDAMENTAIS\n1. A crise moral (Niilismo) exige ação.\n2. A vida deve ser afirmada (não negada).\n3. Autossuperação e criação de valores."
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_COLOR
    
    sh10_3 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(3.9), Inches(4), Inches(3.1))
    sh10_3.fill.solid(); sh10_3.fill.fore_color.rgb = RGBColor(30, 40, 60)
    p = sh10_3.text_frame.paragraphs[0]
    p.text = "5 CONCEITOS-CHAVE\n• Morte de Deus\n• Vontade de Potência\n• Niilismo\n• Amor Fati\n• Übermensch (Além-do-Homem)"
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_COLOR
    
    # Right column quiz
    sh10_4 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.7), Inches(0.5), Inches(8), Inches(6.5))
    sh10_4.fill.solid(); sh10_4.fill.fore_color.rgb = RGBColor(30, 40, 60)
    tf = sh10_4.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Teste sua Compreensão"
    p.font.bold = True; p.font.size = Pt(24); p.font.color.rgb = TITLE_COLOR
    
    quiz_text = "1. O que Nietzsche diagnosticou com a frase \"Deus está morto\"?\n(A) Um evento biológico e científico.  (B) O fim da religião.\n(C) O colapso dos fundamentos absolutos e morais.\n\n2. Na filosofia nietzscheana, a \"Vontade de Potência\" é:\n(A) O desejo egoísta de dominar. (B) A força vital de autossuperação.\n\n3. O que caracteriza o conceito de \"Amor Fati\"?\n(A) Amar apenas alegrias. (B) Amar o destino, abraçando alegremente a vida, inclusive a dor.\n\n4. Qual é a principal crítica de Nietzsche à \"Moral de Rebanho\"?\n(A) Ela nega a vida real e glorifica a submissão. (B) Ela exige força extrema.\n\n5. A metáfora do \"Leão\" representa:\n(A) A aceitação passiva. (B) O espírito de revolta que diz \"Eu quero\"."
    p = tf.add_paragraph()
    p.text = quiz_text
    p.font.size = Pt(14); p.font.color.rgb = TEXT_COLOR
    
    sh10_g = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.7), Inches(6.2), Inches(8), Inches(0.5))
    sh10_g.fill.solid(); sh10_g.fill.fore_color.rgb = RGBColor(22, 30, 45)
    sh10_g.line.color.rgb = TITLE_COLOR
    p = sh10_g.text_frame.paragraphs[0]
    p.text = "GABARITO: 1-C | 2-B | 3-B | 4-A | 5-B"
    p.font.size = Pt(14); p.font.color.rgb = TITLE_COLOR
    p.alignment = PP_ALIGN.CENTER

    # 11. SOURCES
    s11 = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s11)
    add_title(s11, "Image Sources")
    tb11 = s11.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(5))
    tf11 = tb11.text_frame
    p = tf11.paragraphs[0]
    p.text = "As imagens usadas na apresentação original pertencem a seus respectivos criadores e fontes:\n\n• Retrato de Nietzsche: etsy.com\n• Caminhante sobre o mar de névoa (Caspar David Friedrich): mydailyartdisplay.uk\n• Imagem da montanha (IA): craiyon.com"
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_COLOR

    # Salvar pptx
    ppt_file = os.path.join(os.getcwd(), "Apresentacao_Nietzsche.pptx")
    prs.save(ppt_file)
    print(f"PPTX salvo em {ppt_file}")
    
    # Converter para PDF
    pdf_file = os.path.join(os.getcwd(), "Apresentacao_Nietzsche.pdf")
    try:
        powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
        # powerpoint.Visible = 1
        slides = powerpoint.Presentations.Open(ppt_file, WithWindow=False)
        slides.SaveAs(pdf_file, 32) # 32 = ppSaveAsPDF
        slides.Close()
        powerpoint.Quit()
        print(f"PDF salvo em {pdf_file}")
    except Exception as e:
        print(f"Erro ao converter PDF: {e}")

if __name__ == "__main__":
    main()
